# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_student_report():
    doc = Document()

    # 1. Căn lề trang chuẩn (Trái 3.0cm, Phải 2.0cm, Trên 2.0cm, Dưới 2.0cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.79)   # 2.0 cm

    # 2. Font mặc định: Times New Roman, 13pt, Màu đen (#000000)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = normal_style._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rPr.append(rFonts)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=13, space_before=0, space_after=3, line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = RGBColor(0, 0, 0)
            rPr = run._element.get_or_add_rPr()
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            rFonts.set(qn('w:cs'), 'Times New Roman')
            rPr.append(rFonts)
        return p

    def add_run(p, text, bold=False, italic=False, size=13):
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(0, 0, 0)
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')
        rPr.append(rFonts)
        return run

    def set_cell_borders(cell):
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(tcBorders)

    # ==========================================
    # TIÊU ĐỀ BÁO CÁO (ĐÚNG THEO MẪU USER ĐÃ CHỈNH)
    # ==========================================
    add_p("ĐỀ TÀI: CÁCH MẠNG THÁNG TÁM NĂM 1945 DƯỚI SỰ LÃNH ĐẠO CỦA ĐẢNG", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_before=0, space_after=2)
    add_p("Trọng tâm: Cuộc khởi nghĩa giành chính quyền tại Hà Nội", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, size=13, space_after=10)

    # ==========================================
    # PHẦN 1: THÔNG TIN NHÓM VÀ ĐƯỜNG DẪN SẢN PHẨM
    # ==========================================
    add_p("1. THÔNG TIN NHÓM VÀ ĐƯỜNG DẪN SẢN PHẨM", bold=True, size=13, space_before=4, space_after=3)

    p_group = add_p()
    add_run(p_group, "• Nhóm thực hiện: ", bold=True, size=13)
    add_run(p_group, "Nhóm 09", size=13)

    p_class = add_p()
    add_run(p_class, "• Lớp học phần: ", bold=True, size=13)
    add_run(p_class, "LLCT220514_06UTExMC", size=13)

    p_genre = add_p()
    add_run(p_genre, "• Thể loại sản phẩm: ", bold=True, size=13)
    add_run(p_genre, "Game tương tác dạng Visual Novel", size=13)

    p_link = add_p()
    add_run(p_link, "• Link trải nghiệm trực tiếp: ", bold=True, size=13)
    add_run(p_link, "https://shinn1603.github.io/LSD/", size=13)

    p_repo = add_p()
    add_run(p_repo, "• Link mã nguồn: ", bold=True, size=13)
    add_run(p_repo, "https://github.com/shinn1603/LSD", size=13)

    add_p("Bảng thành viên Nhóm 09:", bold=True, italic=True, size=13, space_before=4, space_after=3)

    # Bảng thành viên 3 cột đúng mẫu
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.8), Inches(3.2), Inches(1.8)]
    headers = ["STT", "Họ và Tên", "MSSV"]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_borders(hdr_cells[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_run(p, title, bold=True, size=12)

    # Data Rows
    team_data = [
        ("1", "Nguyễn Phước Minh Triết", "24110357"),
        ("2", "Nguyễn Phước Thọ", "24110343"),
        ("3", "Trần Tiến Đạt", "24110198"),
        ("4", "Lương Xuân Phúc", "24142302"),
        ("5", "Đinh Thiên Bảo", "24162009")
    ]

    for row_idx, row_data in enumerate(team_data, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_borders(row_cells[col_idx])
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            p.alignment = align
            add_run(p, text, bold=(col_idx == 1 and row_idx == 1), size=12)

    add_p("", space_after=4)

    # ==========================================
    # PHẦN 2: HÌNH ẢNH TỔNG THỂ DỰ ÁN
    # ==========================================
    add_p("2. HÌNH ẢNH TỔNG THỂ DỰ ÁN", bold=True, size=13, space_before=6, space_after=4)

    screenshots = [
        ("docs/screenshots/1_title_screen.png", "Hình 1: Màn hình chính"),
        ("docs/screenshots/2_syllabus_modal.png", "Hình 2: Bảng đề cương môn học"),
        ("docs/screenshots/3_codex_modal.png", "Hình 3: Hồ sơ tư liệu lịch sử"),
        ("docs/screenshots/4_quiz_modal.png", "Hình 4: Phân hệ trắc nghiệm kiến thức"),
        ("docs/screenshots/5_gameplay_vn.png", "Hình 5: Giao diện Visual Novel"),
        ("docs/screenshots/6_choices_system.png", "Hình 6: Hệ thống lựa chọn rẽ nhánh"),
        ("docs/screenshots/7_cutscene_opera.png", "Hình 7: Phân cảnh Nhà hát Lớn chiều 17/8/1945"),
        ("docs/screenshots/8_ending_true.png", "Hình 8: Màn hình kết thúc toàn thắng (True Ending)")
    ]

    for img_rel_path, caption in screenshots:
        full_img_path = os.path.join(os.getcwd(), img_rel_path)
        if os.path.exists(full_img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(full_img_path, width=Inches(5.7))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(6)
            add_run(p_cap, caption, bold=True, italic=True, size=11)

    # ==========================================
    # PHẦN 3: HƯỚNG DẪN TRẢI NGHIỆM
    # ==========================================
    add_p("3. HƯỚNG DẪN TRẢI NGHIỆM", bold=True, size=13, space_before=6, space_after=3)

    p_guide = add_p()
    add_run(p_guide, "• Cách chơi: Mở trình duyệt vào link https://shinn1603.github.io/LSD/, bấm chuột hoặc phím Space / Enter để đọc thoại. Bấm phím 'A' để bật chế độ tự động chạy thoại.\n", size=13)
    add_run(p_guide, "• Mã kiểm tra nhanh (Cheat code): Vào phần Cài đặt nhập mã 'yain' (hoặc gõ liên tiếp các phím y-a-i-n bất cứ lúc nào) để mở khóa toàn bộ hồ sơ tư liệu và các nhánh kết thúc", size=13)

    # Lưu file đúng tên mẫu mà user yêu cầu
    output_filename = "261LLCT220514_06MOOC_GameLichSuDang_Nhom09.docx"
    doc.save(output_filename)
    print(f"SUCCESS: Generated {output_filename}")

if __name__ == "__main__":
    create_student_report()
